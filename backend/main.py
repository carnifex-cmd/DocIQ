from __future__ import annotations

import hashlib
import json
import re
import uuid
from pathlib import Path
from typing import Any

import chromadb
import fitz
from chromadb.config import Settings
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from config import EMBEDDING_PROVIDER_CHANGED_MESSAGE, FALLBACK_ANSWER, settings
from providers.embeddings import embed_query, embed_texts
from providers.llm import generate_answer


CHROMA_DIR = Path(__file__).resolve().parent / "chroma_db"
STORE_METADATA_PATH = CHROMA_DIR / "embedding_metadata.json"
INDEX_MANIFEST_PATH = CHROMA_DIR / "indexed_files.json"
COLLECTION_NAME = "business_documents"
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
CHUNK_SIZE = 800
CHUNK_OVERLAP = 150

CHROMA_DIR.mkdir(parents=True, exist_ok=True)
settings.documents_dir.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="Business Document RAG API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.frontend_origins or ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

chroma_client = chromadb.PersistentClient(
    path=str(CHROMA_DIR),
    settings=Settings(
        anonymized_telemetry=False,
        chroma_product_telemetry_impl="chroma_telemetry.NoOpTelemetry",
    ),
)
collection = chroma_client.get_or_create_collection(name=COLLECTION_NAME)


class AskRequest(BaseModel):
    question: str


class AskResponse(BaseModel):
    answer: str
    sources: list[dict[str, Any]]


class IndexErrorResponse(BaseModel):
    source_path: str
    error: str


class IndexResponse(BaseModel):
    total_files_found: int
    files_indexed: int
    files_skipped: int
    chunks_created: int
    errors: list[IndexErrorResponse]


class IndexStatusResponse(BaseModel):
    indexed_files_count: int
    total_chunks: int
    embedding_provider: str
    llm_provider: str


def normalize_text(text: str) -> str:
    cleaned_lines = []
    for line in text.splitlines():
        compact_line = re.sub(r"\s+", " ", line).strip()
        if compact_line:
            cleaned_lines.append(compact_line)
    return "\n".join(cleaned_lines)


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    if not text:
        return []

    chunks: list[str] = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= text_length:
            break
        start = max(end - overlap, 0)

    return chunks


def build_file_metadata(
    source_file: str,
    source_path: str,
    file_type: str,
    file_hash: str,
    modified_time: str,
    page_number: int | None = None,
) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "source_file": source_file,
        "source_path": source_path,
        "file_type": file_type,
        "file_hash": file_hash,
        "modified_time": modified_time,
    }
    if page_number is not None:
        metadata["page_number"] = page_number
    return metadata


def extract_pdf_documents(
    file_path: Path,
    source_file: str,
    source_path: str,
    file_hash: str,
    modified_time: str,
) -> list[dict[str, Any]]:
    documents: list[dict[str, Any]] = []
    pdf = fitz.open(file_path)

    try:
        for page_index, page in enumerate(pdf, start=1):
            text = normalize_text(page.get_text("text"))
            for chunk in chunk_text(text):
                documents.append(
                    {
                        "text": chunk,
                        "metadata": build_file_metadata(
                            source_file=source_file,
                            source_path=source_path,
                            file_type="pdf",
                            file_hash=file_hash,
                            modified_time=modified_time,
                            page_number=page_index,
                        ),
                    }
                )
    finally:
        pdf.close()

    return documents


def extract_docx_text(file_path: Path) -> str:
    document = Document(str(file_path))
    sections: list[str] = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()
        if paragraph_text:
            sections.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                sections.append(" | ".join(cells))

    return normalize_text("\n".join(sections))


def extract_docx_documents(
    file_path: Path,
    source_file: str,
    source_path: str,
    file_hash: str,
    modified_time: str,
) -> list[dict[str, Any]]:
    text = extract_docx_text(file_path)
    documents: list[dict[str, Any]] = []

    for chunk in chunk_text(text):
        documents.append(
            {
                "text": chunk,
                "metadata": build_file_metadata(
                    source_file=source_file,
                    source_path=source_path,
                    file_type="docx",
                    file_hash=file_hash,
                    modified_time=modified_time,
                ),
            }
        )

    return documents


def extract_documents(
    file_path: Path,
    source_file: str,
    source_path: str,
    file_hash: str,
    modified_time: str,
) -> list[dict[str, Any]]:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf_documents(file_path, source_file, source_path, file_hash, modified_time)
    if suffix == ".docx":
        return extract_docx_documents(file_path, source_file, source_path, file_hash, modified_time)

    raise ValueError("Only PDF and DOCX files are supported.")


def current_store_metadata(dimension: int) -> dict[str, Any]:
    return {
        "embedding_provider": settings.embedding_provider,
        "embedding_model": settings.embedding_model,
        "embedding_dimension": dimension,
    }


def read_store_metadata() -> dict[str, Any] | None:
    if not STORE_METADATA_PATH.exists():
        return None

    try:
        return json.loads(STORE_METADATA_PATH.read_text())
    except json.JSONDecodeError:
        return None


def write_store_metadata(dimension: int) -> None:
    STORE_METADATA_PATH.write_text(json.dumps(current_store_metadata(dimension), indent=2))


def ensure_store_compatibility(dimension: int) -> None:
    if collection.count() == 0:
        return

    store_metadata = read_store_metadata()
    if not store_metadata:
        raise HTTPException(status_code=400, detail=EMBEDDING_PROVIDER_CHANGED_MESSAGE)

    expected_metadata = current_store_metadata(dimension)
    for key, value in expected_metadata.items():
        if store_metadata.get(key) != value:
            raise HTTPException(status_code=400, detail=EMBEDDING_PROVIDER_CHANGED_MESSAGE)


def is_dimension_mismatch_error(error: Exception) -> bool:
    message = str(error).lower()
    return "dimension" in message or "embedding" in message


def read_index_manifest() -> dict[str, dict[str, Any]]:
    if not INDEX_MANIFEST_PATH.exists():
        return {}

    try:
        payload = json.loads(INDEX_MANIFEST_PATH.read_text())
    except json.JSONDecodeError:
        return {}

    files = payload.get("files", {})
    if not isinstance(files, dict):
        return {}

    manifest: dict[str, dict[str, Any]] = {}
    for source_path, metadata in files.items():
        if isinstance(source_path, str) and isinstance(metadata, dict):
            manifest[source_path] = metadata
    return manifest


def write_index_manifest(manifest: dict[str, dict[str, Any]]) -> None:
    INDEX_MANIFEST_PATH.write_text(json.dumps({"files": manifest}, indent=2, sort_keys=True))


def reset_collection_for_reindex() -> None:
    global collection

    try:
        chroma_client.delete_collection(name=COLLECTION_NAME)
    except Exception:
        pass

    collection = chroma_client.get_or_create_collection(name=COLLECTION_NAME)

    if STORE_METADATA_PATH.exists():
        STORE_METADATA_PATH.unlink()
    if INDEX_MANIFEST_PATH.exists():
        INDEX_MANIFEST_PATH.unlink()


def rebuild_manifest_from_collection() -> dict[str, dict[str, Any]]:
    if collection.count() == 0:
        return {}

    results = collection.get(include=["metadatas"])
    metadatas = results.get("metadatas", [])
    manifest: dict[str, dict[str, Any]] = {}

    for metadata in metadatas:
        if not isinstance(metadata, dict):
            return {}

        source_path = metadata.get("source_path")
        source_file = metadata.get("source_file")
        file_type = metadata.get("file_type")
        file_hash = metadata.get("file_hash")
        modified_time = metadata.get("modified_time")
        if not all(isinstance(value, str) for value in [source_path, source_file, file_type, file_hash, modified_time]):
            return {}

        entry = manifest.setdefault(
            source_path,
            {
                "source_file": source_file,
                "file_type": file_type,
                "file_hash": file_hash,
                "modified_time": modified_time,
                "chunk_count": 0,
            },
        )
        entry["chunk_count"] += 1

    return manifest


def ensure_index_manifest_ready() -> dict[str, dict[str, Any]]:
    manifest = read_index_manifest()
    if manifest:
        return manifest

    if collection.count() == 0:
        return {}

    rebuilt_manifest = rebuild_manifest_from_collection()
    if rebuilt_manifest:
        write_index_manifest(rebuilt_manifest)
        return rebuilt_manifest

    reset_collection_for_reindex()
    return {}


def scan_documents_folder() -> tuple[list[Path], list[Path], int]:
    supported_files: list[Path] = []
    unsupported_files: list[Path] = []

    if not settings.documents_dir.exists():
        return supported_files, unsupported_files, 0

    for path in sorted(settings.documents_dir.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() in SUPPORTED_EXTENSIONS:
            supported_files.append(path)
        else:
            unsupported_files.append(path)

    total_files_found = len(supported_files) + len(unsupported_files)
    return supported_files, unsupported_files, total_files_found


def to_source_path(file_path: Path) -> str:
    return file_path.relative_to(settings.documents_dir).as_posix()


def compute_file_hash(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as file_handle:
        for chunk in iter(lambda: file_handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def modified_time_for_file(file_path: Path) -> str:
    return str(file_path.stat().st_mtime_ns)


def is_file_already_indexed(
    manifest: dict[str, dict[str, Any]],
    source_path: str,
    modified_time: str,
    file_hash: str,
) -> bool:
    entry = manifest.get(source_path)
    if not entry:
        return False
    return entry.get("modified_time") == modified_time and entry.get("file_hash") == file_hash


def remove_existing_chunks_for_file(source_path: str) -> None:
    try:
        collection.delete(where={"source_path": source_path})
    except Exception:
        pass


def upsert_documents(documents: list[dict[str, Any]]) -> int:
    if not documents:
        return 0

    texts = [item["text"] for item in documents]
    metadatas = [item["metadata"] for item in documents]
    embeddings = embed_texts(texts)

    if not embeddings:
        return 0

    ensure_store_compatibility(len(embeddings[0]))
    ids = [f"chunk-{uuid.uuid4().hex}" for _ in documents]

    try:
        collection.upsert(
            ids=ids,
            documents=texts,
            metadatas=metadatas,
            embeddings=embeddings,
        )
    except Exception as error:
        if is_dimension_mismatch_error(error):
            raise HTTPException(status_code=400, detail=EMBEDDING_PROVIDER_CHANGED_MESSAGE) from error
        raise

    write_store_metadata(len(embeddings[0]))
    return len(documents)


def index_file(file_path: Path, source_path: str, file_hash: str, modified_time: str) -> tuple[int, dict[str, Any]]:
    documents = extract_documents(
        file_path=file_path,
        source_file=file_path.name,
        source_path=source_path,
        file_hash=file_hash,
        modified_time=modified_time,
    )
    if not documents:
        raise ValueError("No extractable text found.")

    chunk_count = upsert_documents(documents)
    if chunk_count == 0:
        raise ValueError("No chunks were created from the document.")

    manifest_entry = {
        "source_file": file_path.name,
        "file_type": file_path.suffix.lower().lstrip("."),
        "file_hash": file_hash,
        "modified_time": modified_time,
        "chunk_count": chunk_count,
    }
    return chunk_count, manifest_entry


def format_context(documents: list[str], metadatas: list[dict[str, Any]]) -> str:
    blocks: list[str] = []

    for index, (document, metadata) in enumerate(zip(documents, metadatas), start=1):
        page_number = metadata.get("page_number")
        page_label = page_number if isinstance(page_number, int) and page_number > 0 else "N/A"
        blocks.append(
            "\n".join(
                [
                    f"Source {index}",
                    f"File: {metadata.get('source_file', 'Unknown')}",
                    f"Path: {metadata.get('source_path', 'Unknown')}",
                    f"Page: {page_label}",
                    "Content:",
                    document,
                ]
            )
        )

    return "\n\n".join(blocks)


def build_sources(metadatas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unique_sources: list[dict[str, Any]] = []
    seen: set[tuple[str, Any]] = set()

    for metadata in metadatas:
        source_file = metadata.get("source_file", "Unknown")
        source_path = metadata.get("source_path", source_file)
        page_number = metadata.get("page_number")
        source_key = (source_path, page_number)
        if source_key in seen:
            continue
        seen.add(source_key)
        unique_sources.append(
            {
                "source_file": source_file,
                "source_path": source_path,
                "page_number": page_number if isinstance(page_number, int) and page_number > 0 else None,
            }
        )

    return unique_sources


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/index/status", response_model=IndexStatusResponse)
def index_status() -> IndexStatusResponse:
    manifest = ensure_index_manifest_ready()
    return IndexStatusResponse(
        indexed_files_count=len(manifest),
        total_chunks=collection.count(),
        embedding_provider=settings.embedding_provider,
        llm_provider=settings.llm_provider,
    )


@app.post("/index", response_model=IndexResponse)
def index_documents() -> IndexResponse:
    settings.documents_dir.mkdir(parents=True, exist_ok=True)
    manifest = ensure_index_manifest_ready()

    supported_files, unsupported_files, total_files_found = scan_documents_folder()
    supported_source_paths = {to_source_path(file_path) for file_path in supported_files}

    removed_paths = sorted(set(manifest) - supported_source_paths)
    for source_path in removed_paths:
        remove_existing_chunks_for_file(source_path)
        manifest.pop(source_path, None)

    files_indexed = 0
    files_skipped = len(unsupported_files)
    chunks_created = 0
    errors: list[IndexErrorResponse] = []

    for file_path in supported_files:
        source_path = to_source_path(file_path)

        try:
            modified_time = modified_time_for_file(file_path)
            file_hash = compute_file_hash(file_path)

            if is_file_already_indexed(manifest, source_path, modified_time, file_hash):
                files_skipped += 1
                continue

            if source_path in manifest:
                remove_existing_chunks_for_file(source_path)

            chunk_count, manifest_entry = index_file(
                file_path=file_path,
                source_path=source_path,
                file_hash=file_hash,
                modified_time=modified_time,
            )
            manifest[source_path] = manifest_entry
            files_indexed += 1
            chunks_created += chunk_count
        except HTTPException:
            raise
        except Exception as error:
            manifest.pop(source_path, None)
            files_skipped += 1
            errors.append(IndexErrorResponse(source_path=source_path, error=str(error)))

    write_index_manifest(manifest)

    return IndexResponse(
        total_files_found=total_files_found,
        files_indexed=files_indexed,
        files_skipped=files_skipped,
        chunks_created=chunks_created,
        errors=errors,
    )


@app.post("/ask", response_model=AskResponse)
def ask_question(payload: AskRequest) -> AskResponse:
    question = payload.question.strip()
    if not question:
        raise HTTPException(status_code=400, detail="Question cannot be empty.")

    if collection.count() == 0:
        return AskResponse(answer=FALLBACK_ANSWER, sources=[])

    query_embedding = embed_query(question)
    ensure_store_compatibility(len(query_embedding))

    try:
        results = collection.query(
            query_embeddings=[query_embedding],
            n_results=5,
            include=["documents", "metadatas", "distances"],
        )
    except Exception as error:
        if is_dimension_mismatch_error(error):
            raise HTTPException(status_code=400, detail=EMBEDDING_PROVIDER_CHANGED_MESSAGE) from error
        raise

    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]

    if not documents or not metadatas:
        return AskResponse(answer=FALLBACK_ANSWER, sources=[])

    context = format_context(documents, metadatas)
    answer = generate_answer(question=question, context=context)
    return AskResponse(answer=answer, sources=build_sources(metadatas))
